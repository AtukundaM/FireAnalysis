import os
import tempfile
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


FOLDER = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(FOLDER, "Fire_Engineering_Annual_Analysis.pptx")
NAVY = "0B2E5C"
BLUE = "#1F4E96"
GOLD = "#C9A227"
RED = "#B22222"
COLORS = {"Fire": BLUE, "Engineering": GOLD}


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.25), Inches(12.2), Inches(0.65))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = RGBColor.from_string(NAVY)
    if subtitle:
        box = slide.shapes.add_textbox(Inches(0.58), Inches(0.88), Inches(11.8), Inches(0.35))
        p = box.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(11)
        p.font.color.rgb = RGBColor(90, 90, 90)


def add_notes(slide, text):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(6.55), Inches(11.9), Inches(0.55))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(70, 70, 70)


def add_bullets(slide, items, left=0.75, top=1.35, width=11.5, height=4.9, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(45, 45, 45)
        p.space_after = Pt(12)
        p.level = 0


def save_chart(fig, temp_dir, name):
    path = os.path.join(temp_dir, name)
    fig.savefig(path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def add_chart(slide, path, left=0.55, top=1.15, width=12.2, height=5.25):
    slide.shapes.add_picture(path, Inches(left), Inches(top), width=Inches(width), height=Inches(height))


def main():
    df = pd.read_excel(os.path.join(FOLDER, "data.xlsx"), sheet_name="Sheet1")
    df = df[df["pol_uw_year"] >= 2012].copy()
    df["prm_gross_written"] = df["prm_gross_written"].fillna(0)
    df["clm_gross_incurred"] = df["clm_gross_incurred"].fillna(0)
    df["claim_policy"] = df["clm_gross_incurred"] > 0
    groups = [g for g in ["Fire", "Engineering"] if g in df["treaty_group_name"].unique()]
    annual = (df.groupby(["treaty_group_name", "pol_uw_year"])
              .agg(policies=("pol_no", "count"), claim_policies=("claim_policy", "sum"),
                   premium=("prm_gross_written", "sum"), claims=("clm_gross_incurred", "sum"))
              .reset_index())
    annual["frequency"] = annual["claim_policies"] / annual["policies"]
    annual["loss_ratio"] = annual["claims"] / annual["premium"].replace(0, np.nan)
    annual["severity"] = annual["claims"] / annual["claim_policies"].replace(0, np.nan)
    totals = annual.groupby("treaty_group_name")["premium", "claims"].sum() if False else df.groupby("treaty_group_name").agg(premium=("prm_gross_written", "sum"), claims=("clm_gross_incurred", "sum"), policies=("pol_no", "count"), claim_policies=("claim_policy", "sum"))
    totals["loss_ratio"] = totals["claims"] / totals["premium"]
    totals["frequency"] = totals["claim_policies"] / totals["policies"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    with tempfile.TemporaryDirectory() as temp_dir:
        slide = prs.slides.add_slide(blank)
        add_title(slide, "Fire and Engineering Risk Analysis", "Annual experience from 2012 onward | Gross basis")
        add_bullets(slide, [
            "Purpose: compare Fire and Engineering performance separately by underwriting year.",
            "Premium measure: prm_gross_written throughout the analysis.",
            "Claim frequency: policies with positive gross incurred claims divided by policies.",
            "Loss ratio: gross incurred claims divided by gross written premium.",
        ], size=20)
        add_notes(slide, "All figures are gross of reinsurance; recent years may have immature claims development.")

        slide = prs.slides.add_slide(blank)
        add_title(slide, "Data quality: classification fields", "Undefined client type is treated as missing or unusable")
        missing_markers = {"", "undefined", "null", "none", "nan"}
        def missing_rate(field):
            values = df[field].fillna("").astype(str).str.strip().str.lower()
            return values.isin(missing_markers).mean()
        client_missing = missing_rate("clt_type")
        construction_missing = missing_rate("ri_construction_class_name")
        add_bullets(slide, [
            f"Client type: {client_missing:.1%} missing, blank, or undefined.",
            f"Construction class: {construction_missing:.1%} missing, blank, or undefined.",
            "An 'undefined' value is not a valid client classification and should be corrected at source or excluded from segmentation.",
            "Completeness should be monitored separately for Fire and Engineering before using these fields for underwriting decisions.",
        ], size=17)
        add_notes(slide, "Data-quality findings affect how confidently portfolio results can be segmented by client type or construction class.")

        fig, ax = plt.subplots(figsize=(10, 4.4))
        rates = pd.DataFrame({"Client type": [missing_rate("clt_type") * 100], "Construction class": [construction_missing * 100]})
        ax.bar(rates.columns, rates.iloc[0], color=[BLUE, GOLD], width=0.55)
        ax.set_ylabel("Missing / undefined records (%)")
        ax.set_title("Classification Data Quality")
        ax.set_ylim(0, 100)
        for i, value in enumerate(rates.iloc[0]):
            ax.text(i, value + 2, f"{value:.1f}%", ha="center")
        ax.spines[["top", "right"]].set_visible(False)
        quality_chart = save_chart(fig, temp_dir, "quality.png")
        slide = prs.slides.add_slide(blank)
        add_title(slide, "Data quality visual", "Higher bars indicate less usable classification data")
        add_chart(slide, quality_chart, top=1.2, height=5.1)
        add_notes(slide, "The client-type series counts 'undefined' as a quality issue, not as a meaningful category.")

        fig, ax = plt.subplots(figsize=(10, 4.5))
        for group in groups:
            view = annual[annual.treaty_group_name == group]
            ax.plot(view.pol_uw_year, view.premium / 1e9, marker="o", linewidth=2, label=group, color=COLORS[group])
        ax.set_title("Annual Gross Written Premium")
        ax.set_xlabel("Underwriting year")
        ax.set_ylabel("Premium (UGX billions)")
        ax.legend(frameon=False)
        ax.spines[["top", "right"]].set_visible(False)
        premium_chart = save_chart(fig, temp_dir, "premium.png")
        slide = prs.slides.add_slide(blank)
        add_title(slide, "Portfolio scale", "Annual gross written premium by treaty group")
        add_chart(slide, premium_chart)
        add_notes(slide, "This chart shows the relative size of the Fire and Engineering books and how that changes over time.")

        fig, axes = plt.subplots(1, 2, figsize=(11, 4.2), sharey=True)
        for ax, group in zip(axes, groups):
            view = annual[annual.treaty_group_name == group]
            ax.plot(view.frequency * 100, view.loss_ratio * 100, "o", color=COLORS[group], markersize=8)
            for _, row in view.iterrows():
                ax.annotate(str(int(row.pol_uw_year)), (row.frequency * 100, row.loss_ratio * 100), xytext=(4, 4), textcoords="offset points", fontsize=8)
            ax.axhline(100, color=RED, linestyle="--", linewidth=1)
            ax.set_title(group)
            ax.set_xlabel("Claim frequency (%)")
            ax.set_ylabel("Gross loss ratio (%)")
            ax.spines[["top", "right"]].set_visible(False)
        fig.suptitle("Frequency versus Loss Ratio by Year", fontsize=13)
        freq_lr_chart = save_chart(fig, temp_dir, "frequency_loss_ratio.png")
        slide = prs.slides.add_slide(blank)
        add_title(slide, "Frequency versus loss ratio", "Each point is one underwriting year")
        add_chart(slide, freq_lr_chart, top=1.1, height=5.45)
        add_notes(slide, "A rising loss ratio with stable frequency usually points toward larger claim severity or premium adequacy; rising frequency indicates broader claims activity.")

        fig, axes = plt.subplots(1, 3, figsize=(13, 4.2))
        for group in groups:
            view = annual[annual.treaty_group_name == group]
            axes[0].plot(view.pol_uw_year, view.claims / 1e9, marker="o", label=group, color=COLORS[group])
            axes[1].plot(view.pol_uw_year, view.loss_ratio * 100, marker="o", label=group, color=COLORS[group])
            axes[2].plot(view.pol_uw_year, view.severity / 1e6, marker="o", label=group, color=COLORS[group])
        for ax, title, ylabel in zip(axes, ["Incurred claims", "Loss ratio", "Average severity"], ["UGX billions", "%", "UGX millions"]):
            ax.set_title(title)
            ax.set_xlabel("Year")
            ax.set_ylabel(ylabel)
            ax.spines[["top", "right"]].set_visible(False)
            ax.tick_params(axis="x", rotation=45)
        axes[1].axhline(100, color=RED, linestyle="--", linewidth=1)
        axes[0].legend(frameon=False)
        fig.suptitle("Annual Claims Outcomes", fontsize=13)
        outcomes_chart = save_chart(fig, temp_dir, "outcomes.png")
        slide = prs.slides.add_slide(blank)
        add_title(slide, "Annual claims outcomes", "Claims, loss ratio, and average severity")
        add_chart(slide, outcomes_chart, top=1.1, height=5.45)
        add_notes(slide, "Loss ratio combines frequency, severity, and premium adequacy. Read the three panels together rather than relying on one metric.")

        slide = prs.slides.add_slide(blank)
        add_title(slide, "Key conclusions and actions")
        add_bullets(slide, [
            "Review Fire and Engineering separately: their exposure mix and loss mechanisms are different.",
            "Use frequency to identify whether claims are becoming more widespread across policies.",
            "Use average severity to identify whether large losses are driving the annual loss ratio.",
            "Treat clt_type = 'undefined' as a data-quality exception and improve capture at point of sale.",
            "Validate construction class completeness before using it for risk selection, pricing, or accumulation analysis.",
            "Interpret the latest years cautiously because claims may still be developing.",
        ], size=17)
        add_notes(slide, "Recommended next step: agree data-ownership actions for client type and construction class, then refresh the annual monitoring pack.")

    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()